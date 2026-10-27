"""Text out of a teacher's material: PDF, Word (.docx), PowerPoint (.pptx) or plain text.

Only the text is read: scanned pages (images) give nothing, and the caller reports that.
"""

import io
import zipfile

from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_FILE_BYTES = 10 * 1024 * 1024
# What is sent to the model at most: enough for a chapter or a lecture, and a bounded bill.
MAX_CHARS = 60_000

KINDS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".txt": "txt",
}


class UnreadableMaterial(Exception):
    """The file is broken, of another kind than its name says, or has no text."""


def kind_of(filename: str) -> str | None:
    name = filename.lower()
    return next((kind for ext, kind in KINDS.items() if name.endswith(ext)), None)


def _pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, ValueError, KeyError, TypeError) as error:
        raise UnreadableMaterial from error


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as error:
        raise UnreadableMaterial from error
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    try:
        presentation = Presentation(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as error:
        raise UnreadableMaterial from error
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


def _txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1251", errors="replace")


READERS = {"pdf": _pdf, "docx": _docx, "pptx": _pptx, "txt": _txt}


def text_of(kind: str, data: bytes) -> str:
    """The material's text, whitespace tidied, cut to MAX_CHARS. Raises UnreadableMaterial if none."""
    text = READERS[kind](data)
    lines = [" ".join(line.split()) for line in text.splitlines()]
    tidy = "\n".join(line for line in lines if line)
    if len(tidy) < 40:
        raise UnreadableMaterial
    return tidy[:MAX_CHARS]

"""Questions from teaching material. The model is never called here: `generate` is replaced."""

import io

import pytest

from app.ai import materials, questions
from app.config import get_settings

LESSON = "Фотосинтез — бул өсүмдүктөрдүн жарыктын жардамы менен органикалык зат түзүүсү. " * 3


@pytest.fixture
def fake_ai(monkeypatch):
    calls = []

    async def generate(material, count, language):
        calls.append((material, count, language))
        return "Фотосинтез", [{"text": "Эмне?", "answers": ["a", "b", "", ""], "correct": 2}]

    monkeypatch.setattr(questions, "generate", generate)
    monkeypatch.setattr(get_settings(), "openai_api_key", "test-key")
    return calls


def _docx(text: str) -> bytes:
    from docx import Document

    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def _pptx(text: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    buffer = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = text
    presentation.save(buffer)
    return buffer.getvalue()


def _pdf(text: str) -> bytes:
    """A one-page PDF with a line of ASCII text, written by hand."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for n, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % n + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return out


def test_text_comes_out_of_every_kind_of_file():
    assert "Фотосинтез" in materials.text_of("docx", _docx(LESSON))
    assert "Фотосинтез" in materials.text_of("pptx", _pptx(LESSON))
    assert "Photosynthesis" in materials.text_of(
        "pdf", _pdf("Photosynthesis makes sugar from light and water.")
    )
    assert materials.text_of("txt", ("Урок про клетки. " * 5).encode("cp1251")).startswith("Урок")
    with pytest.raises(materials.UnreadableMaterial):
        materials.text_of("docx", b"not a zip")
    with pytest.raises(materials.UnreadableMaterial):
        materials.text_of("txt", b"too short")
    assert len(materials.text_of("txt", ("x" * 100 + "\n").encode() * 1000)) == materials.MAX_CHARS
    assert materials.kind_of("Lesson.PDF") == "pdf" and materials.kind_of("a.doc") is None


def test_model_output_is_checked_again():
    raw = {
        "title": " Тема ",
        "questions": [
            {"text": "Q1", "answers": ["a", "b", "c", "d"], "correct": 4},
            {"text": "Q2", "answers": ["a", "b"], "correct": 1},
            {"text": "", "answers": ["a", "b"], "correct": 1},
            {"text": "Q4", "answers": ["a"], "correct": 1},
            {"text": "Q5", "answers": ["a", "b"], "correct": 3},
            {"text": "Q6", "answers": ["a", " "], "correct": 1},
            {"text": "Q7", "answers": ["a", "b", "c", "d", "e"], "correct": 1},
        ],
    }
    title, items = questions._clean(raw, 10)
    assert title == "Тема"
    assert items == [
        {"text": "Q1", "answers": ["a", "b", "c", "d"], "correct": 4},
        {"text": "Q2", "answers": ["a", "b", "", ""], "correct": 1},
    ]
    assert len(questions._clean(raw, 1)[1]) == 1


async def test_a_teacher_gets_questions_from_a_file_or_text(as_user, fake_ai):
    c, _ = await as_user("ulan")
    files = {"file": ("lesson.docx", _docx(LESSON), "application/octet-stream")}
    r = await c.post("/ai/questions", data={"count": "5", "language": "ky"}, files=files)
    assert r.status_code == 200, r.json()
    assert r.json() == {
        "title": "Фотосинтез",
        "questions": [{"text": "Эмне?", "answers": ["a", "b", "", ""], "correct": 2}],
    }
    r = await c.post("/ai/questions", data={"count": "10", "language": "ru", "text": LESSON})
    assert r.status_code == 200
    assert [(n, lang) for _, n, lang in fake_ai] == [(5, "ky"), (10, "ru")]


async def test_bad_requests_get_codes(as_user, fake_ai):
    c, _ = await as_user("ulan")
    form = {"count": "5", "language": "ky"}
    assert (await c.post("/ai/questions", data=form)).json()["code"] == "MATERIAL_REQUIRED"
    r = await c.post("/ai/questions", data=form, files={"file": ("a.doc", b"x" * 100, "application/msword")})
    assert (r.status_code, r.json()["code"]) == (415, "UNSUPPORTED_FILE")
    r = await c.post("/ai/questions", data=form, files={"file": ("a.pdf", b"%PDF broken", "application/pdf")})
    assert r.json()["code"] == "MATERIAL_UNREADABLE"
    assert (await c.post("/ai/questions", data={**form, "text": "short"})).json()[
        "code"
    ] == "MATERIAL_TOO_SHORT"
    for bad in (
        {"count": "0", "language": "ky"},
        {"count": "31", "language": "ky"},
        {"count": "5", "language": "en"},
    ):
        assert (await c.post("/ai/questions", data={**bad, "text": LESSON})).status_code == 422
    assert fake_ai == []


async def test_only_teachers_and_a_daily_limit(as_user, fake_ai, monkeypatch):
    admin, _ = await as_user("root", role="admin")
    form = {"count": "5", "language": "ky", "text": LESSON}
    assert (await admin.post("/ai/questions", data=form)).status_code == 403
    teacher, _ = await as_user("ulan")
    monkeypatch.setattr(get_settings(), "ai_daily_limit", 2)
    codes = [(await teacher.post("/ai/questions", data=form)).status_code for _ in range(3)]
    assert codes == [200, 200, 429]


async def test_no_key_no_feature(as_user, fake_ai, monkeypatch):
    monkeypatch.setattr(get_settings(), "openai_api_key", "")
    teacher, _ = await as_user("ulan")
    r = await teacher.post("/ai/questions", data={"count": "5", "language": "ky", "text": LESSON})
    assert (r.status_code, r.json()["code"]) == (503, "AI_NOT_CONFIGURED")
